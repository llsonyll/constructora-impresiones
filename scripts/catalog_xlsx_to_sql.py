"""Genera SQL de importación (proveedores, productos, costos) desde el Excel de catálogo.

Uso: python3 scripts/catalog_xlsx_to_sql.py catalogo.xlsx [lote] > import.sql   (requiere openpyxl)

Idempotente: proveedores por nombre, productos por SKU (actualiza nombre/precio/etc.).
Productos sin precio se cargan inactivos (la BD no permite activos sin precio).
No toca el stock: se maneja por stock_movements.
"""
import json, sys
import openpyxl

def sheet_rows(ws):
    header = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            yield dict(zip(header, r))

def clean(v):
    return v.strip() if isinstance(v, str) and v.strip() else (None if isinstance(v, str) else v)

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)

providers = [{k: clean(v) for k, v in p.items()} for p in sheet_rows(wb['Proveedores'])]
products = []
for p in sheet_rows(wb['Productos']):
    p = {k: clean(v) for k, v in p.items()}
    notes = '; '.join(x for x in [
        f"Precio: {p['price_source']}" if p.get('price_source') else None,
        p.get('notas'),
        f"Proveedor: {p['provider']}" if p.get('provider') and p['provider'] not in {q['name'] for q in providers} else None,
    ] if x)
    products.append({
        'sku': p['sku'], 'name': p['name'], 'brand': p.get('marca'), 'category': p['category'],
        'unit': p.get('unit') or 'und', 'price': p.get('price'), 'cost': p.get('cost'),
        'provider': p.get('provider'), 'notes': notes or None,
        'needs_review': str(p.get('revisar') or '').upper() == 'SI',
    })

def lit(obj):
    obj = [{k: v for k, v in o.items() if v is not None} for o in obj]  # jsonb_to_recordset rellena con null
    return "'" + json.dumps(obj, ensure_ascii=False, separators=(',', ':')).replace("'", "''") + "'::jsonb"

BATCH = int(sys.argv[2]) if len(sys.argv) > 2 else 0   # 0 = todo en un bloque
chunks = [products[i:i + BATCH] for i in range(0, len(products), BATCH)] if BATCH else [products]

out = [f"""insert into providers (name, ruc, contact, phone, email, address, notes)
select x.name, x.ruc, x.contact, x.phone, x.email, x.address, x.notes
from jsonb_to_recordset({lit(providers)}) as x(name text, ruc text, contact text, phone text, email text, address text, notes text)
where not exists (select 1 from providers p where p.name = x.name);"""]

for chunk in chunks:
    out.append(f"""with x as (
  select * from jsonb_to_recordset({lit(chunk)}) as x(sku text, name text, brand text, category text, unit text,
       price numeric, cost numeric, provider text, notes text, needs_review boolean)
), up as (
  insert into products (sku, name, brand, category_id, provider_id, unit, price, active, notes, needs_review)
  select x.sku, x.name, x.brand, c.id, pr.id, x.unit, x.price, x.price is not null, x.notes, x.needs_review
  from x
  left join categories c on c.name = x.category
  left join providers pr on pr.name = x.provider
  on conflict (sku) do update set
    name = excluded.name, brand = excluded.brand, category_id = excluded.category_id,
    provider_id = excluded.provider_id, unit = excluded.unit, price = excluded.price,
    active = excluded.active, notes = excluded.notes, needs_review = excluded.needs_review
  returning id, sku
)
insert into product_costs (product_id, cost)
select up.id, x.cost from up join x on x.sku = up.sku
where x.cost is not null
on conflict (product_id) do update set cost = excluded.cost, updated_at = now();""")

# Con BATCH, cada bloque va separado por una línea "-- @@batch" para ejecutarlos por partes.
print("\n-- @@batch\n".join(out) if BATCH else "begin;\n" + "\n\n".join(out) + "\ncommit;")
